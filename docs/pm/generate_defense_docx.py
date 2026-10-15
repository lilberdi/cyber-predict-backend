#!/usr/bin/env python3
"""Word report for the course defense."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).with_name("07-Otchet_na_zashchitu.docx")


def font(run, size=14, bold=False, italic=False, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def p(doc, text, *, align="justify", indent=True, size=14, bold=False, italic=False, after=6):
    para = doc.add_paragraph()
    para.alignment = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
    }[align]
    pf = para.paragraph_format
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.5
    pf.first_line_indent = Cm(1.25) if indent else Cm(0)
    run = para.add_run(text)
    font(run, size=size, bold=bold, italic=italic)
    return para


def h(doc, text):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_before = Pt(12)
    para.paragraph_format.space_after = Pt(8)
    para.paragraph_format.first_line_indent = Cm(0)
    run = para.add_run(text)
    font(run, size=16, bold=True)


def bullets(doc, items):
    for item in items:
        para = doc.add_paragraph(style="List Bullet")
        para.clear()
        para.paragraph_format.left_indent = Cm(1.25)
        para.paragraph_format.first_line_indent = Cm(0)
        para.paragraph_format.line_spacing = 1.15
        run = para.add_run(item)
        font(run, size=14)


def table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, title in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(title)
        font(run, size=11, bold=True, color=RGBColor(255, 255, 255))
        tcPr = cell._tc.get_or_add_tcPr()
        shd = tcPr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): "1F4E79"})
        tcPr.append(shd)
    for r_i, row in enumerate(rows):
        fill = "F2F2F2" if r_i % 2 else "FFFFFF"
        for c_i, val in enumerate(row):
            cell = t.rows[r_i + 1].cells[c_i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(val)
            font(run, size=10)
            tcPr = cell._tc.get_or_add_tcPr()
            shd = tcPr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): fill})
            tcPr.append(shd)
    doc.add_paragraph()


def build() -> None:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(3)
        section.right_margin = Cm(1.5)

    p(doc, "Дисциплина: «Управление проектами»", align="center", indent=False, bold=True, after=0)
    p(doc, "ОТЧЁТ ПО ПРОЕКТУ ДЛЯ ЗАЩИТЫ", align="center", indent=False, bold=True, size=18, after=8)
    p(
        doc,
        "Разработка интеллектуальной системы прогнозирования результатов киберспортивных матчей (CS2). Cyber Predict.",
        align="center",
        indent=False,
        bold=True,
        size=14,
        after=8,
    )
    p(
        doc,
        "Команда: 4 человека. Горизонт: 12 недель. Дата: 25 сентября 2026. Статус: недели 1–4 + Swagger + живой календарь PandaScore.",
        align="center",
        indent=False,
        italic=True,
    )
    p(
        doc,
        "Прогнозы учебные, не ставки и не финансовая рекомендация. HLTV.org не используется как источник данных.",
        italic=True,
    )

    h(doc, "1. Цель")
    p(
        doc,
        "Прототип, который берёт календарь предстоящих матчей CS2 из легального источника, оценивает вероятность победы по pre-match признакам, отдаёт JSON через REST и документирует контракт в Swagger. Сопровождается пакетом документов PM.",
    )

    h(doc, "2. Управление проектом")
    p(
        doc,
        "Объём ограничен методом MoSCoW. Must: календарь, вероятность, метрики, Docker, документы, легальные данные. Should: XGBoost, UI, калибровка. Won’t: ставки, аккаунты, парсинг HLTV, другие дисциплины. Решение связано со сроком (12 недель) и составом (4 человека).",
    )
    table(
        doc,
        ["Спринт", "Недели", "Содержание", "Статус"],
        [
            ["S1", "1–2", "Устав, каркас Go, схема БД", "сделано"],
            ["S2", "3–4", "История, Elo+logreg, OpenAPI", "сделано"],
            ["+", "после S2", "Swagger, hourly PandaScore, топ-60", "сделано"],
            ["S3", "5–6", "XGBoost, тир, состав", "не начато"],
            ["S4", "7–8", "Frontend", "не начато"],
            ["S5–S6", "9–12", "Калибровка, отчёт, защита", "этот документ"],
        ],
    )

    h(doc, "3. Данные и право")
    p(
        doc,
        "HLTV.org не парсим: в Terms of Service (п. 2.2) запрещены data mining, scraping и конкурирующий сервис. Официального API нет.",
    )
    table(
        doc,
        ["Контур", "Источник"],
        [
            ["Обучение", "Sample 480 матчей; позже licensed датасет"],
            ["Календарь", "PandaScore GET /csgo/matches/upcoming (CS2, title=13)"],
            ["Фолбэк", "data/fixtures/upcoming.json"],
        ],
    )
    p(
        doc,
        "Календарь обновляется эндпоинтом GET/POST /api/v1/jobs/sync-upcoming и тикером раз в час. Фильтр топ-набора (~60): по умолчанию достаточно одной известной команды. Два ноунейма отбрасываются. Это не рейтинг HLTV, а свой allow-list плюс Elo проекта.",
    )

    h(doc, "4. Архитектура")
    p(
        doc,
        "API — Go/Gin. Обучение — Python (Elo + логистическая регрессия). Граница — PostgreSQL. В обычном GET модель не вызывается: Gin читает готовые predictions. Sync сам пишет календарь и быстрый Elo-прогноз (elo_sync).",
    )
    table(
        doc,
        ["Слой", "Стек"],
        [
            ["API", "Go, Gin, pgx, golang-migrate, Docker Compose"],
            ["ML", "Python, pandas, scikit-learn, joblib"],
            ["Документация", "OpenAPI 3, Swagger UI /swagger/"],
        ],
    )

    h(doc, "5. Реализованный API")
    bullets(
        doc,
        [
            "GET /health — живость и Postgres;",
            "GET /api/v1/matches/upcoming — ближайшие матчи и прогноз;",
            "GET /api/v1/matches/{id}, /predictions/{id} — карточка и вероятности;",
            "GET /api/v1/model/metrics — accuracy, log-loss, Brier, сравнение с Elo;",
            "GET/POST /api/v1/jobs/sync-upcoming — обновление календаря;",
            "GET /swagger/ — интерактивная спека по разделам.",
        ],
    )

    h(doc, "6. Критерии прогноза и качество")
    p(
        doc,
        "Цель: вероятность победы team1 в серии. Признаки только из прошлого: elo_diff, form5_diff, form10_diff, h2h_diff. Не используем счёт текущего матча и post-match стату (утечка). StandardScaler + LogisticRegression. Confidence: low / medium / high по отрыву от 0,5.",
    )
    table(
        doc,
        ["Метрика", "Значение"],
        [
            ["Accuracy (тест 96 матчей)", "≈ 65,6%"],
            ["Elo-фаворит без ML", "≈ 65,6%"],
            ["Log-loss", "≈ 0,66"],
            ["Brier", "≈ 0,23"],
        ],
    )
    p(
        doc,
        "Для CS2 55–65% — нормальный baseline. Модель пока не обгоняет правило «всегда фаворит Elo»; это честно фиксируем. XGBoost — спринт S3. Проценты в Swagger — не live-нейросеть, а заранее записанные строки в БД.",
    )

    h(doc, "7. Сценарий демонстрации")
    bullets(
        doc,
        [
            "Показать устав и этот отчёт (HLTV нельзя, MoSCoW);",
            "Открыть /swagger/, GET /health;",
            "GET /matches/upcoming — source pandascore, вероятности;",
            "Разобрать Elo и confidence на одном матче;",
            "GET /model/metrics;",
            "GET /jobs/sync-upcoming — fetched / kept / dropped;",
            "Токен PandaScore на экран не выводить.",
        ],
    )

    h(doc, "8. Риски")
    table(
        doc,
        ["Риск", "Реакция"],
        [
            ["Free PandaScore без глубокой статы", "Календарь с API, модель на истории"],
            ["Утечка признаков", "Только окна до даты матча"],
            ["Нет сети на защите", "Docker + fixtures"],
            ["Претензия HLTV", "Запрет scrape"],
            ["Модель не лучше Elo", "Честное сравнение, elo_sync как fallback"],
        ],
    )

    h(doc, "9. Что дальше")
    p(
        doc,
        "XGBoost; licensed датасет; признаки тира и состава; frontend (список, деталка, факт vs прогноз); калибровка и итоговая презентация.",
    )

    h(doc, "10. Источники")
    p(
        doc,
        "HLTV.org Terms of Service, п. 2.2. PandaScore Developers, CS2: префикс /csgo/, videogame_title=13. Репозиторий проекта и пакет docs/pm.",
        indent=False,
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
