#!/usr/bin/env python3
"""Build a course Word file from the week 1–4 PM pack."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).with_name("Plan_proekta_prognozirovanie_CS2.docx")


def font(run, size=14, bold=False, italic=False, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def p(doc, text, *, align="justify", indent=True, size=14, bold=False, italic=False, after=6):
    para = doc.add_paragraph()
    para.alignment = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
    }[align]
    pf = para.paragraph_format
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.5
    pf.first_line_indent = Cm(1.25) if indent else Cm(0)
    run = para.add_run(text)
    font(run, size=size, bold=bold, italic=italic)
    return para


def h(doc, text, level=1):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_before = Pt(12)
    para.paragraph_format.space_after = Pt(8)
    para.paragraph_format.first_line_indent = Cm(0)
    run = para.add_run(text)
    font(run, size=16 if level == 1 else 14, bold=True)
    return para


def bullets(doc, items):
    for item in items:
        para = doc.add_paragraph(style="List Bullet")
        para.clear()
        para.paragraph_format.left_indent = Cm(1.25)
        para.paragraph_format.first_line_indent = Cm(0)
        para.paragraph_format.line_spacing = 1.15
        run = para.add_run(item)
        font(run, size=14)


def table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, title in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(title)
        font(run, size=11, bold=True, color=RGBColor(255, 255, 255))
        tcPr = cell._tc.get_or_add_tcPr()
        shd = tcPr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): "1F4E79"})
        tcPr.append(shd)
    for r_i, row in enumerate(rows):
        fill = "F2F2F2" if r_i % 2 else "FFFFFF"
        for c_i, val in enumerate(row):
            cell = t.rows[r_i + 1].cells[c_i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(val)
            font(run, size=11)
            tcPr = cell._tc.get_or_add_tcPr()
            shd = tcPr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): fill})
            tcPr.append(shd)
    doc.add_paragraph()
    return t


def build() -> None:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(3)
        section.right_margin = Cm(1.5)

    p(doc, "Дисциплина: «Управление проектами»", align="center", indent=False, bold=True, after=0)
    p(doc, "ПРОЕКТНАЯ РАБОТА", align="center", indent=False, bold=True, size=18, after=12)
    p(
        doc,
        "Разработка интеллектуальной системы прогнозирования результатов киберспортивных матчей (Counter-Strike 2)",
        align="center",
        indent=False,
        bold=True,
        size=16,
        after=18,
    )
    p(doc, "Устав, план работ, стек и отчёт по неделям 1–4", align="center", indent=False, italic=True, after=24)
    p(doc, "Команда: 4 человека. Срок: 12 недель. Документ: 2026.", align="center", indent=False)

    h(doc, "1. Устав проекта")
    p(
        doc,
        "Название: Cyber Predict — интеллектуальная система прогнозирования результатов профессиональных матчей Counter-Strike 2. Цель: воспроизводимый прототип, который для предстоящего матча оценивает вероятность победы команды и отдаёт её через REST API (далее — в веб-интерфейсе).",
    )
    p(doc, "Критерии успеха:", indent=True)
    bullets(
        doc,
        [
            "прогноз на реальные upcoming или на фикстуры защиты;",
            "метрики модели и сравнение с Elo-фаворитом;",
            "обоснованный отказ от парсинга HLTV;",
            "пакет документов PM;",
            "демо через Docker Compose.",
        ],
    )
    p(
        doc,
        "Ограничения: только CS2; без ставок, аккаунтов и мобильного приложения; HLTV.org не используется как источник (Terms of Service, п. 2.2: запрещены data mining, web scraping и создание похожего сервиса).",
    )

    h(doc, "2. Стек и архитектура")
    p(
        doc,
        "Публичный API — Go и Gin. Обучение — Python (к неделе 4: Elo + логистическая регрессия; XGBoost запланирован на недели 5–6). Единственная граница между языками — PostgreSQL: Python пишет матчи, метрики и прогнозы, Gin только читает.",
    )
    table(
        doc,
        ["Слой", "Технологии"],
        [
            ["API", "Go, Gin, pgx, golang-migrate, Docker Compose"],
            ["ML", "Python 3.11, pandas, scikit-learn, joblib"],
            ["БД", "PostgreSQL 16"],
            ["Контракт", "docs/openapi.yaml"],
            ["Frontend (позже)", "Next.js, TypeScript, Tailwind"],
        ],
    )
    p(
        doc,
        "История для обучения: синтетический sample (недели 3–4) с возможностью подменить licensed датасетом Kaggle/Hugging Face. Upcoming: PandaScore GET /csgo/matches/upcoming?filter[videogame_title]=13 либо data/fixtures/upcoming.json.",
    )

    h(doc, "3. Стейкхолдеры и RACI")
    table(
        doc,
        ["Стейкхолдер", "Интерес"],
        [
            ["Преподаватель", "Документы PM, сроки, прозрачные решения"],
            ["Команда (4 человека)", "Работающий прототип и сдача"],
            ["Пользователь-фан CS2", "Список матчей и вероятность"],
            ["PandaScore / авторы датасета", "Соблюдение лицензии, без scrape HLTV"],
        ],
    )
    table(
        doc,
        ["Работа", "PM", "ML", "Go", "UI"],
        [
            ["Документы PM", "A/R", "C", "C", "I"],
            ["Данные и лицензии", "A", "R", "C", "I"],
            ["Схема БД и API", "C", "I", "A/R", "C"],
            ["Обучение модели", "I", "A/R", "I", "I"],
            ["Frontend", "C", "I", "C", "A/R"],
            ["Защита", "A/R", "C", "C", "C"],
        ],
    )

    h(doc, "4. MoSCoW, WBS и график")
    p(doc, "Must: upcoming с вероятностью; карточка матча; факт vs прогноз; страница модели; Docker; документы PM. Should: фильтры, калибровка, переобучение. Won’t: ставки, аккаунты, другие дисциплины, парсинг HLTV.")
    table(
        doc,
        ["Спринт", "Недели", "Содержание", "Статус"],
        [
            ["S1", "1–2", "Устав, каркас Go, схема БД, клиент PandaScore", "сделано"],
            ["S2", "3–4", "История, EDA, Elo+logreg, OpenAPI", "сделано"],
            ["S3", "5–6", "XGBoost, live-прогноз", "не начато"],
            ["S4", "7–8", "UI", "не начато"],
            ["S5", "9–10", "Репетиция защиты", "Compose уже есть"],
            ["S6", "11–12", "Отчёт и защита", "не начато"],
        ],
    )

    h(doc, "5. Реестр рисков")
    table(
        doc,
        ["ID", "Риск", "Реакция"],
        [
            ["R1", "Free PandaScore без глубокой статы", "Модель на истории; API = календарь"],
            ["R2", "Утечка post-match признаков", "Только окна до даты матча + тест"],
            ["R3", "Нет сети на защите", "Docker + fixtures"],
            ["R4", "Два языка", "Контракт только через Postgres"],
            ["R5", "Претензия HLTV", "Запрет scrape в уставе и ревью"],
        ],
    )

    h(doc, "6. Статус на конец недели 4")
    p(
        doc,
        "Выполнено: пакет PM; API /health, /api/v1/matches/upcoming, /matches/:id, /predictions/:id, /model/metrics; миграции и Compose; OpenAPI; Python-конвейер generate-sample → load-history → train → sync → predict. Ключ PandaScore опционален. Следующий спринт: XGBoost и расширение признаков.",
    )
    p(doc, "Источники: HLTV.org Terms of Service; PandaScore developers.pandascore.co; настоящий репозиторий.", italic=True)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
