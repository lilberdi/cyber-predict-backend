#!/usr/bin/env python3
"""Word copy of 06-progress-and-criteria.md for the course."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).with_name("06-Otchet_progress_i_kriterii_prognozov.docx")


def font(run, size=14, bold=False, italic=False, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def p(doc, text, *, align="justify", indent=True, size=14, bold=False, italic=False, after=6):
    para = doc.add_paragraph()
    para.alignment = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
    }[align]
    pf = para.paragraph_format
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.5
    pf.first_line_indent = Cm(1.25) if indent else Cm(0)
    run = para.add_run(text)
    font(run, size=size, bold=bold, italic=italic)
    return para


def h(doc, text, level=1):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_before = Pt(12)
    para.paragraph_format.space_after = Pt(8)
    para.paragraph_format.first_line_indent = Cm(0)
    run = para.add_run(text)
    font(run, size=16 if level == 1 else 14, bold=True)
    return para


def bullets(doc, items):
    for item in items:
        para = doc.add_paragraph(style="List Bullet")
        para.clear()
        para.paragraph_format.left_indent = Cm(1.25)
        para.paragraph_format.first_line_indent = Cm(0)
        para.paragraph_format.line_spacing = 1.15
        run = para.add_run(item)
        font(run, size=14)


def table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, title in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(title)
        font(run, size=11, bold=True, color=RGBColor(255, 255, 255))
        tcPr = cell._tc.get_or_add_tcPr()
        shd = tcPr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): "1F4E79"})
        tcPr.append(shd)
    for r_i, row in enumerate(rows):
        fill = "F2F2F2" if r_i % 2 else "FFFFFF"
        for c_i, val in enumerate(row):
            cell = t.rows[r_i + 1].cells[c_i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(val)
            font(run, size=10)
            tcPr = cell._tc.get_or_add_tcPr()
            shd = tcPr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): fill})
            tcPr.append(shd)
    doc.add_paragraph()
    return t


def build() -> None:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(3)
        section.right_margin = Cm(1.5)

    p(doc, "Дисциплина: «Управление проектами»", align="center", indent=False, bold=True, after=0)
    p(doc, "ОТЧЁТ О ХОДЕ ПРОЕКТА И КРИТЕРИЯХ ПРОГНОЗОВ", align="center", indent=False, bold=True, size=16, after=8)
    p(
        doc,
        "Cyber Predict — прогнозирование результатов матчей CS2. Контрольная точка: конец недели 4 + Swagger. 25 сентября 2026.",
        align="center",
        indent=False,
        italic=True,
        after=12,
    )
    p(
        doc,
        "Прогнозы учебные, не ставки. HLTV.org не используется как источник (Terms of Service, п. 2.2).",
        italic=True,
    )

    h(doc, "1. Что уже сделано")
    p(doc, "Управление проектом: устав, стейкхолдеры, RACI, MoSCoW, WBS, спринт-план, реестр рисков.")
    p(doc, "Backend (Go/Gin) читает только PostgreSQL. Модель в HTTP-запросе не вызывается.")
    bullets(
        doc,
        [
            "GET /health — живость сервиса и связь с БД;",
            "GET /api/v1/matches/upcoming — ближайшие матчи и вложенный прогноз;",
            "GET /api/v1/matches/{id} — карточка матча;",
            "GET /api/v1/predictions/{id} — вероятность, уверенность, топ-признаки;",
            "GET /api/v1/model/metrics — метрики обучения;",
            "Swagger UI /swagger/, Docker Compose, порт из .env.",
        ],
    )
    p(
        doc,
        "Данные и модель: 480 sample-матчей; календарь PandaScore или фикстуры; Elo + логистическая регрессия; признаки только из прошлого; python -m ml predict пишет таблицу predictions.",
    )
    table(
        doc,
        ["Метрика", "Значение", "Смысл"],
        [
            ["Accuracy", "≈ 65,6%", "Доля верных исходов на тесте (96 матчей)"],
            ["Elo-фаворит", "≈ 65,6%", "Без ML: всегда побеждает больший Elo"],
            ["Log-loss", "≈ 0,66", "Качество вероятностей (меньше — лучше)"],
            ["Brier", "≈ 0,23", "То же для калибровки"],
        ],
    )
    p(
        doc,
        "Для CS2 55–65% — нормальный baseline. Модель пока не обгоняет правило «фаворит Elo»; это честный результат недели 4.",
    )

    h(doc, "2. Что предстоит")
    table(
        doc,
        ["Спринт", "Недели", "Работы", "Статус"],
        [
            ["S1", "1–2", "Устав, каркас Go, схема БД", "сделано"],
            ["S2", "3–4", "История, Elo+logreg, Swagger", "сделано"],
            ["S3", "5–6", "XGBoost, тир турнира, состав", "не начато"],
            ["S4", "7–8", "Frontend: список, деталка, факт vs прогноз", "не начато"],
            ["S5", "9–10", "Калибровка, офлайн-демо, репетиция", "Compose уже есть"],
            ["S6", "11–12", "Итоговый отчёт и защита", "не начато"],
        ],
    )
    p(doc, "Won’t: ставки, аккаунты, live по раундам, другие дисциплины, парсинг HLTV.")

    h(doc, "3. Почему в API уже есть предикты")
    p(
        doc,
        "Это не нейросеть в момент открытия Swagger. Прогнозы заранее посчитаны офлайн и лежат в PostgreSQL. Gin только читает таблицу predictions.",
    )
    bullets(
        doc,
        [
            "load-history — 480 завершённых матчей (sample), по ним считается сила команд;",
            "train — Elo и логистическая регрессия, метрики в model_metrics;",
            "sync — 8 предстоящих встреч (фикстуры или PandaScore); названия команд реальные, расписание демонстрационное;",
            "predict — для каждой пары берёт текущие Elo, форму и H2H только из уже сыгранных матчей;",
            "GET /matches/upcoming джойнит predictions — поэтому вероятности видны сразу.",
        ],
    )
    p(
        doc,
        "Предикты есть, потому что контур недель 3–4 уже прогнан. Это демонстрация архитектуры «модель пишет, API читает», а не live-инференс и не букмекерская линия. Обновление календаря — снова sync и predict. Переобучение — только если меняется история или модель.",
    )

    h(doc, "4. По каким критериям считаются предикты")
    p(
        doc,
        "Цель: вероятность победы team1 в серии (BO1/BO3/BO5 как один исход). Фаворит: p_team1_win ≥ 0,5. Признаки — только из матчей строго раньше прогнозируемого. Счёт текущего матча не используется.",
    )
    table(
        doc,
        ["Признак", "Смысл", "Как считается"],
        [
            ["elo_diff", "Разница силы", "Elo старт 1500, K=32; R1 − R2"],
            ["form5_diff", "Недавняя форма", "Винрейт за 5 матчей: team1 минус team2"],
            ["form10_diff", "Длиннее форма", "То же за 10 матчей"],
            ["h2h_diff", "Личные встречи", "Доля побед team1 за 12 мес. минус 0,5"],
        ],
    )
    p(doc, "Не используются: счёт и карты текущего матча; ADR/K/D/рейтинг после игры (утечка); veto, тир и состав — спринт S3.")
    p(doc, "Как получается число в API:", indent=True)
    bullets(
        doc,
        [
            "четыре признака масштабируются (StandardScaler);",
            "логистическая регрессия даёт P(победа team1); p_team2_win = 1 − p_team1;",
            "confidence: |p−0,5| < 0,07 — low; < 0,15 — medium; иначе high;",
            "в top_features и explanation — вклад Elo, формы и H2H.",
        ],
    )
    p(
        doc,
        "Если Spirit — team1 и p = 0,58, модель оценивает Spirit выше по Elo/форме/H2H из sample-истории, а не «потому что так на HLTV». Сейчас история синтетическая с теми же названиями команд: предикт иллюстрирует пайплайн. После licensed датасета или живого PandaScore те же критерии пойдут по реальным матчам, код признаков менять не нужно.",
    )

    h(doc, "5. Где смотреть")
    p(
        doc,
        "Markdown: docs/pm/06-progress-and-criteria.md. Swagger: http://localhost:<PORT>/swagger/. Метрики: GET /api/v1/model/metrics. Признаки: ml/features.py. Уверенность: ml/predict.py.",
        indent=False,
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
